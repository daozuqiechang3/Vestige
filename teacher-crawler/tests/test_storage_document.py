import csv
from pathlib import Path

from docx import Document
from PIL import Image

from crawler.document import create_teacher_document
from crawler.models import CrawlState, Teacher
from crawler.storage import Storage, teacher_content_hash, teacher_stem


def test_storage_state_csv_and_document(tmp_path: Path) -> None:
    storage = Storage(tmp_path / "output")
    record = Teacher(
        school="Test University",
        college="School of Engineering",
        name="Alice / Zhang",
        category="Doctoral supervisor",
        title="Professor",
        email="alice@example.edu",
        phone="010-12345678",
        department="Computer Science",
        research_interests=["AI", "Robotics"],
        admission_info="Accepting students.",
        sections={"education": "PhD"},
        full_text="First paragraph.\nSecond paragraph.",
        profile_url="https://example.edu/people/alice",
    )

    json_path = storage.save_record(record)
    photo_path = tmp_path / "photo.png"
    Image.new("RGB", (120, 160), "#dddddd").save(photo_path)
    word_path = create_teacher_document(record, storage.documents_dir, photo_path)
    content_hash = teacher_content_hash(record)
    assert teacher_content_hash(
        record.model_copy(update={"profile_url": "https://example.edu/people/alice-copy"})
    ) == content_hash
    storage.state.completed[record.profile_url] = json_path.relative_to(storage.root).as_posix()
    storage.state.content_hashes[content_hash] = record.profile_url
    storage.save_state()
    storage.state.failed["https://example.edu/people/failed"] = "HTTPStatusError: 500"
    failures_path = storage.write_failures()
    storage.state = CrawlState()
    records = storage.load_records()
    csv_path = storage.write_csv(records)

    assert json_path.exists()
    assert json_path.name == "Test University_School of Engineering_Alice _ Zhang.json"
    assert word_path.exists()
    assert word_path.name == "Test University_School of Engineering_Alice _ Zhang.docx"
    assert storage.state_path.exists()
    assert len(records) == 1
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]["name"] == "Alice / Zhang"
    assert rows[0]["college"] == "School of Engineering"
    assert rows[0]["research_interests"] == "AI | Robotics"
    with failures_path.open(encoding="utf-8-sig", newline="") as handle:
        failures = list(csv.DictReader(handle))
    assert failures == [
        {
            "profile_url": "https://example.edu/people/failed",
            "error": "HTTPStatusError: 500",
        }
    ]

    document = Document(word_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    assert paragraph_text[0] == "Alice / Zhang"
    assert "基本信息" in paragraph_text
    assert "研究方向" in paragraph_text
    assert "招生信息" in paragraph_text
    assert "个人主页完整内容" in paragraph_text
    assert "来源与采集" in paragraph_text
    assert len(document.inline_shapes) == 1
    assert [cell.text for cell in document.tables[0].columns[0].cells] == [
        "项目",
        "职称",
        "邮箱",
        "电话",
        "院系",
    ]
    list_items = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Bullet"
    ]
    assert list_items[:2] == ["AI", "Robotics"]
    hyperlink_targets = {
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype.endswith("/hyperlink")
    }
    assert record.profile_url in hyperlink_targets


def test_teacher_stem_replaces_all_windows_forbidden_characters() -> None:
    teacher = Teacher(
        school='A\\B/C:D*E?F"G<H>I|J',
        college="Computer Science",
        name="Alice",
        profile_url="https://example.edu/alice",
    )

    stem = teacher_stem(teacher)

    assert not any(character in stem for character in '\\/ :*?"<>|'.replace(" ", ""))
