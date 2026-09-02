from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .models import Teacher
from .storage import teacher_stem

FONT_NAME = "Microsoft YaHei"
HEADING_BLUE = RGBColor(0x2E, 0x74, 0xB5)
HEADING_DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x66, 0x66)
TABLE_FILL = "E8EEF5"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_COLUMN_WIDTHS_DXA = (2700, 6660)


def _set_font(run: object, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    style_tokens = {
        "Normal": (11, RGBColor(0, 0, 0), 0, 6, 1.25),
        "Heading 1": (16, HEADING_BLUE, 18, 10, 1.0),
        "Heading 2": (13, HEADING_BLUE, 14, 7, 1.0),
        "Heading 3": (12, HEADING_DARK_BLUE, 10, 5, 1.0),
        "List Bullet": (11, RGBColor(0, 0, 0), 0, 4, 1.25),
        "List Number": (11, RGBColor(0, 0, 0), 0, 4, 1.25),
    }
    for style_name, (size, color, before, after, line_spacing) in style_tokens.items():
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line_spacing

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        document.styles[style_name].font.bold = True

    for style_name in ("List Bullet", "List Number"):
        paragraph_format = document.styles[style_name].paragraph_format
        paragraph_format.left_indent = Inches(0.375)
        paragraph_format.first_line_indent = Inches(-0.188)


def _set_cell_margins(cell: object) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table: object) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_columns = table._tbl.tblGrid.gridCol_lst
    for index, width in enumerate(TABLE_COLUMN_WIDTHS_DXA):
        grid_columns[index].set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(TABLE_COLUMN_WIDTHS_DXA[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(TABLE_COLUMN_WIDTHS_DXA[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)


def _shade_cell(cell: object, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _add_hyperlink(paragraph: object, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)
    properties.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([properties, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_page_number(paragraph: object) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    _set_font(run, 9)
    run._r.extend([begin, instruction, end])


def _add_running_furniture(document: Document, teacher: Teacher) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = f"{teacher.school}  |  {teacher.college}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        _set_font(run, 9)
        run.font.color.rgb = MUTED
    _add_page_number(section.footer.paragraphs[0])


def _add_title_block(document: Document, teacher: Teacher) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run(teacher.name)
    _set_font(title_run, 24, bold=True)
    title_run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    subtitle_parts = [teacher.school, teacher.college]
    if teacher.category:
        subtitle_parts.append(teacher.category)
    subtitle = document.add_paragraph(" / ".join(subtitle_parts))
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    for run in subtitle.runs:
        _set_font(run, 11)
        run.font.color.rgb = MUTED


def _add_photo(document: Document, teacher: Teacher, photo_path: Path | None) -> None:
    if not photo_path or not photo_path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    try:
        shape = paragraph.add_run().add_picture(str(photo_path), width=Inches(1.55))
        shape._inline.docPr.set("descr", f"{teacher.name}的教师照片")
        shape._inline.docPr.set("title", f"{teacher.name}照片")
    except (OSError, ValueError):
        paragraph._element.getparent().remove(paragraph._element)


def _add_basic_info(document: Document, teacher: Teacher) -> None:
    document.add_heading("基本信息", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    fields = [
        ("职称", teacher.title),
        ("邮箱", teacher.email),
        ("电话", teacher.phone),
        ("院系", teacher.department),
    ]
    header_cells = table.add_row().cells
    header_cells[0].text = "项目"
    header_cells[1].text = "内容"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_properties.append(header_marker)
    for cell in header_cells:
        _shade_cell(cell, TABLE_FILL)
        for run in cell.paragraphs[0].runs:
            _set_font(run, 10.5, bold=True)
    for label, value in fields:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value or "未提供"
        _shade_cell(cells[0], "F2F4F7")
        for run in cells[0].paragraphs[0].runs:
            _set_font(run, 10.5, bold=True)
        for run in cells[1].paragraphs[0].runs:
            _set_font(run, 10.5)
    _set_table_geometry(table)


def _add_list_section(document: Document, title: str, values: list[str]) -> None:
    document.add_heading(title, level=1)
    if not values:
        document.add_paragraph("未提取到相关信息。")
        return
    for value in values:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(value)


def _add_text_section(document: Document, title: str, content: str | None) -> None:
    document.add_heading(title, level=1)
    if not content:
        document.add_paragraph("未提取到相关信息。")
        return
    for block in content.split("\n"):
        if block.strip():
            document.add_paragraph(block.strip())


def _add_full_text(document: Document, teacher: Teacher) -> None:
    document.add_heading("个人主页完整内容", level=1)
    section_titles = {title for title in teacher.sections if title != "summary"}
    bullet_pattern = re.compile(r"^[*\-•]\s+(.+)$")
    number_pattern = re.compile(r"^(?:\d+|[一二三四五六七八九十]+)[.、]\s*(.+)$")
    for block in teacher.full_text.split("\n"):
        text = block.strip()
        if not text:
            continue
        if text in section_titles:
            document.add_heading(text, level=2)
            continue
        bullet = bullet_pattern.match(text)
        numbered = number_pattern.match(text)
        if bullet:
            document.add_paragraph(bullet.group(1), style="List Bullet")
        elif numbered:
            document.add_paragraph(numbered.group(1), style="List Number")
        else:
            document.add_paragraph(text)


def _add_source(document: Document, teacher: Teacher) -> None:
    document.add_heading("来源与采集", level=1)
    source = document.add_paragraph()
    label = source.add_run("信息来源：")
    _set_font(label, 11, bold=True)
    _add_hyperlink(source, "访问原始个人主页", teacher.profile_url)
    collected = document.add_paragraph()
    label = collected.add_run("采集时间：")
    _set_font(label, 11, bold=True)
    value = collected.add_run(teacher.collected_at.strftime("%Y-%m-%d %H:%M:%S"))
    _set_font(value, 11)


def create_teacher_document(
    record: Teacher, output_dir: Path, photo_path: Path | None = None
) -> Path:
    document = Document()
    _configure_document(document)
    _add_running_furniture(document, record)
    _add_title_block(document, record)
    _add_photo(document, record, photo_path)
    _add_basic_info(document, record)
    _add_list_section(document, "研究方向", record.research_interests)
    _add_text_section(document, "招生信息", record.admission_info)
    _add_full_text(document, record)
    _add_source(document, record)

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{teacher_stem(record)}.docx"
    document.save(path)
    return path
